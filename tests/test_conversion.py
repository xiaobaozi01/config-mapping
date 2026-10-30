"""使用独立 cfg/xlsx fixture 验证端到端转换流程。"""

from __future__ import annotations

import json
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from config_adaptor.adaptation.service import convert
from config_adaptor.common.errors import InvariantViolation


FIXTURES = Path(__file__).parent / "fixtures"


def conversion_fixture(name: str) -> tuple[Path, Path]:
    """返回端到端场景的拓扑文件和设备配置目录。"""
    root = FIXTURES / name
    return root / "topology.xlsx", root / "configs"


class ConversionTest(unittest.TestCase):
    """覆盖两个厂商的 NNI、非 NNI 清理、认证及 group 关键边界。"""

    def test_iosxr_preconfigured_interfaces_are_ignored(self):
        topology, source_config_dir = conversion_fixture("iosxr_bundle")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config_dir = root / "configs"
            shutil.copytree(source_config_dir, config_dir)
            config_path = config_dir / "R1.cfg"
            source = config_path.read_text(encoding="utf-8")
            source = source.replace(
                "\nend\n",
                "\ninterface preconfigure GigabitEthernet 0/0/0/5\n"
                " ipv4 address 198.51.100.1/31\n"
                "!\n"
                "interface preconfigure GigabitEthernet 0/0/0/6\n"
                " description RESERVED\n!\n"
                "interface PTP0/RP0/CPU0/0\n"
                " ipv4 address 203.0.113.1/31\n!\n"
                "banner exec ^\n"
                "interface FutureBanner0\n"
                "= production warning =\n"
                "^\n"
                "ptp\n clock\n!\nend\n",
            )
            config_path.write_text(source, encoding="utf-8")

            output = root / "output"
            context = convert(topology, config_dir, output)
            self.assertFalse(context.has_errors, context.errors)
            converted = (output / "configs" / "R1.cfg").read_text(encoding="utf-8")
            self.assertNotIn("interface preconfigure", converted)
            self.assertNotIn("interface PTP", converted)
            self.assertNotIn("banner exec", converted)
            self.assertNotIn("production warning", converted)
            self.assertNotIn("description RESERVED", converted)
            self.assertIn("ptp\n clock", converted)
            self.assertNotIn("ipv4 address 198.51.100.1/31", converted)
            self.assertFalse(any(
                item.source_interface in {
                    "GigabitEthernet0/0/0/5",
                    "GigabitEthernet0/0/0/6",
                }
                for item in context.devices["R1"].mappings
            ))
            self.assertFalse(any(
                "PTP0/RP0/CPU0/0 类型无法识别" in warning
                for warning in context.devices["R1"].warnings
            ))
            self.assertFalse(any(
                "FutureBanner0" in warning
                for warning in context.devices["R1"].warnings
            ))
            report = json.loads((output / "report.json").read_text(encoding="utf-8"))
            preconfigured_event = next(
                event
                for event in report["events"]
                if event["kind"] == "preconfigured-interface-cleanup"
            )
            self.assertEqual(
                preconfigured_event["removed_by_type"],
                {"preconfigured-interface": 2},
            )
            self.assertTrue(any(
                event["kind"] == "cleaning-rule"
                and event["rule_id"] == "cisco.remove.ptp.interface"
                for event in report["events"]
            ))
            banner_event = next(
                event
                for event in report["events"]
                if event.get("rule_id") == "cisco.remove.banner"
            )
            self.assertEqual(banner_event["point"], "pre_analysis")
            self.assertEqual(banner_event["removed_by_type"], {"banner": 1})
            self.assertEqual(
                report["summary"]["cleaning_rule_hits_by_category"]["ptp-interface"],
                1,
            )

    def test_iosxr_bundle_nni_non_nni_cleanup_and_auth(self):
        topology, config_dir = conversion_fixture("iosxr_bundle")
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output"
            context = convert(topology, config_dir, output)
            self.assertFalse(context.has_errors)
            self.assertNotIn("GigabitEthernet0/0/0/5", context.devices["R1"].profile.nni_interfaces)
            self.assertNotIn("GigabitEthernet0/0/0/6", context.devices["R1"].profile.nni_interfaces)
            self.assertNotIn("GigabitEthernet0/0/0/7", context.devices["R1"].profile.nni_interfaces)
            converted = (output / "configs" / "R1.cfg").read_text(encoding="utf-8")
            self.assertIn("interface Bundle-Ether10", converted)
            self.assertNotIn("old-user", converted)
            self.assertNotIn("tacacs-server", converted)
            self.assertNotIn("taskgroup OLD-TASKS", converted)
            self.assertNotIn("usergroup OLD-USERS", converted)
            self.assertNotIn("login authentication OLD-AUTH", converted)
            self.assertNotIn("authorization commands OLD-AUTHZ", converted)
            self.assertNotIn("accounting commands OLD-ACCT", converted)
            self.assertNotIn("users group OLD-USERS", converted)
            self.assertNotIn("ACCESS-UNI", converted)
            self.assertIn("username root", converted)
            self.assertIn("group root-system", converted)
            self.assertIn("line template vty", converted)
            self.assertIn("exec-timeout 10 0", converted)
            self.assertIn("interface GigabitEthernet0/0/0/0", converted)
            self.assertNotIn("interface GigabitEthernet0/0/0/5", converted)
            self.assertNotIn("interface GigabitEthernet0/0/0/7", converted)
            self.assertNotIn("second-dot1q 2", converted)
            self.assertNotIn("UNUSED-BARE-PORT", converted)
            self.assertIn("interface Loopback0", converted)
            self.assertLess(converted.index("username root"), converted.index("end"))

            adapted = load_workbook(output / "topology-adapted.xlsx")
            adapted_links = adapted["链接表"]
            self.assertEqual(adapted_links.max_row, 2)
            self.assertEqual(adapted_links.cell(2, 2).value, "GigabitEthernet0/0/0/0")
            self.assertEqual(adapted["设备列表"].cell(2, 3).value, "configs/R1.cfg")
            report = json.loads((output / "report.json").read_text(encoding="utf-8"))
            self.assertEqual(report["status"], "success")
            self.assertEqual(report["summary"]["skipped_links"], 1)
            authentication = next(
                event
                for event in report["events"]
                if event.get("rule_id") == "cisco.replace.authentication"
            )
            self.assertEqual(authentication["point"], "post_rewrite")
            self.assertEqual(authentication["added_by_type"], {"lab-account": 1})
            self.assertEqual(authentication["removed_by_type"]["taskgroup"], 1)
            self.assertEqual(authentication["removed_by_type"]["usergroup"], 1)
            self.assertEqual(authentication["removed_by_type"]["line-auth-reference"], 5)
            adjustment = next(
                event for event in report["events"] if event["kind"] == "simulation-adaptation"
            )
            self.assertEqual(adjustment["mode"], "stable")
            self.assertEqual(adjustment["image"], "xrv9000")
            self.assertIn("GigabitEthernet0/0/0/0", adjustment["target_interfaces"])

    def test_unknown_and_virtual_interfaces_are_preserved_but_not_mapped(self):
        """未识别及已知虚拟接口保留原配置，且未知类型生成审计告警。"""
        topology, source_config_dir = conversion_fixture("iosxr_bundle")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config_dir = root / "configs"
            shutil.copytree(source_config_dir, config_dir)
            config_path = config_dir / "R1.cfg"
            source = config_path.read_text(encoding="utf-8")
            source = source.replace(
                "\nend\n",
                "\ninterface Tunnel-ip100\n"
                " ipv4 address 198.51.100.1 255.255.255.252\n"
                "!\n"
                "interface FutureVirtual0\n"
                " ipv4 address 203.0.113.1 255.255.255.252\n"
                "!\n"
                "end\n",
            )
            config_path.write_text(source, encoding="utf-8")

            output = root / "output"
            context = convert(topology, config_dir, output)

            self.assertFalse(context.has_errors)
            converted = (output / "configs" / "R1.cfg").read_text(encoding="utf-8")
            self.assertIn("interface Tunnel-ip100", converted)
            self.assertIn("interface FutureVirtual0", converted)
            self.assertTrue(
                any(
                    "FutureVirtual0 类型无法识别" in warning
                    for warning in context.devices["R1"].warnings
                )
            )
            self.assertFalse(
                any(
                    "Tunnel-ip100 类型无法识别" in warning
                    for warning in context.devices["R1"].warnings
                )
            )
            mapped_sources = {
                mapping.source_interface for mapping in context.devices["R1"].mappings
            }
            self.assertNotIn("Tunnel-ip100", mapped_sources)
            self.assertNotIn("FutureVirtual0", mapped_sources)

    def test_non_physical_nni_endpoint_fails_before_mapping(self):
        """链接表引用未知接口时必须失败，不能猜测为物理口继续转换。"""
        topology, source_config_dir = conversion_fixture("iosxr_bundle")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config_dir = root / "configs"
            shutil.copytree(source_config_dir, config_dir)
            config_path = config_dir / "R1.cfg"
            source = config_path.read_text(encoding="utf-8").replace(
                "\nend\n",
                "\ninterface FutureVirtual0\n"
                " ipv4 address 203.0.113.1 255.255.255.252\n"
                "!\n"
                "end\n",
            )
            config_path.write_text(source, encoding="utf-8")

            invalid_topology = root / "topology.xlsx"
            workbook = load_workbook(topology)
            workbook["链接表"].cell(2, 2).value = "FutureVirtual0"
            workbook.save(invalid_topology)

            output = root / "output"
            context = convert(invalid_topology, config_dir, output)

            self.assertTrue(context.has_errors)
            self.assertTrue(
                any(
                    "FutureVirtual0" in error and "不能作为物理 NNI" in error
                    for error in context.errors
                )
            )
            self.assertFalse((output / "configs").exists())

    def test_junos_bundle_nni_non_nni_cleanup_and_auth(self):
        topology, config_dir = conversion_fixture("junos_bundle")
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output"
            context = convert(topology, config_dir, output)
            self.assertFalse(context.has_errors)
            self.assertNotIn("ge-0/0/5", context.devices["J1"].profile.nni_interfaces)
            self.assertNotIn("ge-0/0/6", context.devices["J1"].profile.nni_interfaces)
            self.assertNotIn("ge-0/0/7", context.devices["J1"].profile.nni_interfaces)
            converted = (output / "configs" / "J1.cfg").read_text(encoding="utf-8")
            self.assertIn("ae0", converted)
            self.assertNotIn("old-user", converted)
            self.assertNotIn("radius-server", converted)
            self.assertNotIn("ACCESS-UNI", converted)
            self.assertNotIn("user root", converted)
            self.assertIn("root-authentication encrypted-password", converted)
            self.assertIn("interface ae0.0;", converted)
            self.assertNotIn("ge-0/0/5", converted)
            self.assertNotIn("ge-0/0/7", converted)
            self.assertNotIn("unit 100", converted)
            self.assertNotIn("vlan-tags outer 100 inner 100;", converted)
            self.assertNotIn("UNUSED-BARE-PORT", converted)
            self.assertIn("lo0", converted)

    def test_junos_non_nni_evpn_vpws_is_removed_with_references(self):
        topology, source_config_dir = conversion_fixture("junos_bundle")
        cases = {
            "leaf": ("unit 0;", ("0",)),
            "block": ("unit 0 { family ccc; }", ("0",)),
            "multiple": (
                "unit 0 { family ccc; } unit 1 { family ccc; }",
                ("0", "1"),
            ),
        }
        for case, (units, numbers) in cases.items():
            with self.subTest(case=case), tempfile.TemporaryDirectory() as directory:
                root = Path(directory)
                config_dir = root / "configs"
                shutil.copytree(source_config_dir, config_dir)
                config_path = config_dir / "J1.cfg"
                source = config_path.read_text(encoding="utf-8")
                source = source.replace(
                    "    lo0 {",
                    "    et-12/0/3 {\n"
                    "        enable;\n"
                    "        mtu 9192;\n"
                    "        encapsulation ethernet-ccc;\n"
                    f"        {units}\n"
                    "    }\n"
                    "    lo0 {",
                )
                for number in numbers:
                    source += (
                        f"\nrouting-instances {{ VPWS_{number} {{\n"
                        "    instance-type evpn-vpws;\n"
                        "    protocols { evpn {\n"
                        f"        interface et-12/0/3.{number} {{\n"
                        "            vpws-service-id { local 999; remote 777; }\n"
                        "        }\n"
                        "    } }\n"
                        f"    interface et-12/0/3.{number};\n"
                        "    route-distinguisher 24208:210101;\n"
                        "    vrf-target target:24208:210101;\n"
                        "} }\n"
                    )
                config_path.write_text(source, encoding="utf-8")

                context = convert(topology, config_dir, root / "output")
                self.assertFalse(context.has_errors, context.errors)
                converted = (root / "output" / "configs" / "J1.cfg").read_text(
                    encoding="utf-8"
                )
                self.assertNotIn("et-12/0/3", converted)
                self.assertNotIn("ge-0/0/5", converted)
                self.assertNotIn("local 999;", converted)
                self.assertNotIn("remote 777;", converted)
                self.assertNotIn("route-distinguisher 24208:210101;", converted)
                self.assertNotIn("vrf-target target:24208:210101;", converted)
                for number in numbers:
                    self.assertNotIn(f"VPWS_{number}", converted)
                self.assertIn("interface ae0.0;", converted)
                self.assertFalse(any(
                    mapping.target_interface and mapping.role != "NNI"
                    for mapping in context.devices["J1"].mappings
                ))

    def test_junos_inactive_children_follow_interface_migration_and_removal(self):
        topology, source_config_dir = conversion_fixture("junos_bundle")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config_dir = root / "configs"
            shutil.copytree(source_config_dir, config_dir)
            config_path = config_dir / "J1.cfg"
            source = config_path.read_text(encoding="utf-8")
            source = source.replace(
                "    ae0 {\n        unit 0 {",
                "    ae0 {\n        inactive: description KEEP-WITH-NNI;\n        unit 0 {",
            ).replace(
                "    ge-0/0/4 {\n        description UNUSED-BARE-PORT;",
                "    ge-0/0/4 {\n        inactive: description DROP-WITH-BARE;\n"
                "        description UNUSED-BARE-PORT;",
            )
            config_path.write_text(source, encoding="utf-8")

            output = root / "output"
            context = convert(topology, config_dir, output)
            self.assertFalse(context.has_errors, context.errors)
            converted = (output / "configs" / "J1.cfg").read_text(encoding="utf-8")
            self.assertIn("inactive: description KEEP-WITH-NNI;", converted)
            self.assertNotIn("DROP-WITH-BARE", converted)
            self.assertIn("ae0", converted)

    def test_cross_vendor_iosxr_to_junos_nni_and_aggregate(self):
        """Cisco 与 Juniper 端点独立适配，并共同删除冗余聚合成员行。"""
        topology, config_dir = conversion_fixture("cross_vendor")
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output"
            context = convert(topology, config_dir, output)
            self.assertFalse(context.has_errors)

            cisco = (output / "configs" / "C1.cfg").read_text(encoding="utf-8")
            self.assertIn("interface Bundle-Ether10", cisco)
            self.assertIn("interface GigabitEthernet0/0/0/0", cisco)
            self.assertIn("ipv4 address 10.0.0.1 255.255.255.252", cisco)
            self.assertIn("interface GigabitEthernet0/0/0/1", cisco)
            self.assertIn("ipv4 address 10.0.1.1 255.255.255.252", cisco)
            self.assertIn("  interface Bundle-Ether10", cisco)
            self.assertIn("  interface GigabitEthernet0/0/0/1", cisco)

            juniper = (output / "configs" / "J1.cfg").read_text(encoding="utf-8")
            self.assertIn("ae20", juniper)
            self.assertIn("ge-0/0/0", juniper)
            self.assertIn("address 10.0.0.2/30;", juniper)
            self.assertIn("ge-0/0/1", juniper)
            self.assertIn("address 10.0.1.2/30;", juniper)
            self.assertIn("interface ae20.0;", juniper)
            self.assertIn("interface ge-0/0/1.0;", juniper)

            adapted = load_workbook(output / "topology-adapted.xlsx")
            links = adapted["链接表"]
            self.assertEqual(links.max_row, 3)
            self.assertEqual(links.cell(2, 2).value, "GigabitEthernet0/0/0/0")
            self.assertEqual(links.cell(2, 4).value, "ge-0/0/0")
            self.assertEqual(links.cell(3, 2).value, "GigabitEthernet0/0/0/1")
            self.assertEqual(links.cell(3, 4).value, "ge-0/0/1")

            mappings = json.loads(
                (output / "interface-mapping.json").read_text(encoding="utf-8")
            )
            self.assertTrue(
                any(
                    item["device"] == "C1"
                    and item["source_interface"] == "Bundle-Ether10"
                    and item["target_interface"] == "Bundle-Ether10"
                    for item in mappings
                )
            )
            self.assertTrue(
                any(
                    item["device"] == "J1"
                    and item["source_interface"] == "ae20"
                    and item["target_interface"] == "ae20"
                    for item in mappings
                )
            )
            report = json.loads((output / "report.json").read_text(encoding="utf-8"))
            self.assertEqual(report["status"], "success")
            self.assertEqual(report["summary"]["active_links"], 2)
            self.assertEqual(report["summary"]["skipped_links"], 1)

    def test_huawei_peer_adapts_supported_ends_and_keeps_huawei_side(self):
        """华为设备不转换，但 Cisco/Juniper 端仍完成 NNI 适配并保留华为链路。"""
        topology, config_dir = conversion_fixture("huawei_peer")
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output"
            context = convert(topology, config_dir, output)
            self.assertFalse(context.has_errors, context.errors)
            self.assertEqual(set(context.devices), {"R1", "J1"})
            self.assertEqual(context.external_devices, {"HW1", "HW2"})

            cisco = (output / "configs" / "R1.cfg").read_text(encoding="utf-8")
            self.assertIn("interface GigabitEthernet0/0/0/0", cisco)
            self.assertIn("ipv4 address 10.0.0.1 255.255.255.252", cisco)
            self.assertIn("interface GigabitEthernet0/0/0/1", cisco)
            self.assertIn("ipv4 address 10.0.1.1 255.255.255.252", cisco)
            self.assertNotIn("GigabitEthernet0/0/0/2", cisco)
            self.assertNotIn("GigabitEthernet0/0/0/3", cisco)

            juniper = (output / "configs" / "J1.cfg").read_text(encoding="utf-8")
            self.assertIn("ge-0/0/0", juniper)
            self.assertIn("address 10.0.0.2/30;", juniper)
            self.assertIn("ge-0/0/1", juniper)
            self.assertIn("address 10.0.1.2/30;", juniper)
            self.assertNotIn("ge-0/0/2", juniper)
            self.assertNotIn("ge-0/0/3", juniper)

            self.assertFalse((output / "configs" / "HW1.cfg").exists())
            self.assertFalse((output / "configs" / "HW2.cfg").exists())

            adapted = load_workbook(output / "topology-adapted.xlsx")
            devices = adapted["设备列表"]
            self.assertEqual(devices.cell(4, 1).value, "HW1")
            self.assertEqual(devices.cell(4, 2).value, "Huawei")
            self.assertIsNone(devices.cell(4, 3).value)
            self.assertEqual(devices.cell(5, 1).value, "HW2")
            self.assertEqual(devices.cell(5, 2).value, "华为")
            self.assertIsNone(devices.cell(5, 3).value)

            links = adapted["链接表"]
            self.assertEqual(links.max_row, 5)
            self.assertEqual(links.cell(2, 2).value, "GigabitEthernet0/0/0/0")
            self.assertEqual(links.cell(2, 4).value, "GigabitEthernet0/0/0")
            self.assertEqual(links.cell(3, 2).value, "ge-0/0/0")
            self.assertEqual(links.cell(3, 4).value, "GigabitEthernet0/0/1")
            self.assertEqual(links.cell(4, 2).value, "GigabitEthernet0/0/0/1")
            self.assertEqual(links.cell(4, 4).value, "ge-0/0/1")
            # 华为与华为之间的链路原样保留，不产生映射。
            self.assertEqual(links.cell(5, 2).value, "GigabitEthernet0/0/2")
            self.assertEqual(links.cell(5, 4).value, "GigabitEthernet0/0/0")

            mappings = json.loads(
                (output / "interface-mapping.json").read_text(encoding="utf-8")
            )
            self.assertFalse(
                any(item["device"] in {"HW1", "HW2"} for item in mappings)
            )

            report = json.loads((output / "report.json").read_text(encoding="utf-8"))
            self.assertEqual(report["status"], "success")
            self.assertEqual(report["summary"]["active_links"], 4)
            self.assertEqual(report["summary"]["skipped_links"], 0)
            self.assertTrue(
                any(
                    "HW1" in warning and "华为" in warning
                    for warning in report["warnings"]
                )
            )
            self.assertTrue(
                any(
                    "HW2" in warning and "华为" in warning
                    for warning in report["warnings"]
                )
            )

    def test_unknown_vendor_device_still_fails_conversion(self):
        """未知厂商仍属于输入错误，不会被当作可忽略的对端。"""
        topology, config_dir = conversion_fixture("iosxr_bundle")
        with tempfile.TemporaryDirectory() as directory:
            temporary_topology = Path(directory) / "topology.xlsx"
            workbook = load_workbook(topology)
            workbook["设备列表"].append(["X1", "H3C", None])
            workbook["链接表"].cell(2, 3).value = "X1"
            workbook["链接表"].cell(2, 4).value = "GigabitEthernet0/0/0"
            workbook.save(temporary_topology)

            output = Path(directory) / "output"
            context = convert(temporary_topology, config_dir, output)
            self.assertTrue(context.has_errors)
            report = json.loads((output / "report.json").read_text(encoding="utf-8"))
            self.assertEqual(report["status"], "failed")
            self.assertTrue(
                any(
                    "X1" in error and "厂商无法识别" in error
                    for error in report["errors"]
                )
            )

    def test_external_peer_multiple_bundle_links_cannot_be_compressed_without_peer_config(self):
        topology, config_dir = conversion_fixture("iosxr_bundle")
        with tempfile.TemporaryDirectory() as directory:
            temporary_topology = Path(directory) / "topology.xlsx"
            workbook = load_workbook(topology)
            for row in range(2, workbook["设备列表"].max_row + 1):
                if workbook["设备列表"].cell(row, 1).value == "R2":
                    workbook["设备列表"].cell(row, 2).value = "Huawei"
                    workbook["设备列表"].cell(row, 3).value = None
            workbook.save(temporary_topology)
            output = Path(directory) / "output"
            context = convert(temporary_topology, config_dir, output)
            self.assertTrue(context.has_errors)
            self.assertEqual(set(context.devices), {"R1"})
            self.assertFalse((output / "configs").exists())
            self.assertTrue(any("外部对端" in error for error in context.errors))
            self.assertIn("interface Bundle-Ether10", context.devices["R1"].document.render())
            self.assertFalse(any(mapping.is_mlag_clone for mapping in context.mappings))

    def test_iosxr_cross_peer_bundle_requires_explicit_mlag_support(self):
        """不同设备不默认视作同一个聚合系统，不能复制逻辑业务。"""
        topology, config_dir = conversion_fixture("iosxr_mlag")
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output"
            context = convert(topology, config_dir, output)
            self.assertTrue(context.has_errors)
            self.assertFalse((output / "configs").exists())
            self.assertFalse((output / "topology-adapted.xlsx").exists())
            self.assertIn("interface Bundle-Ether10.100", context.devices["A"].document.render())
            self.assertFalse(any(mapping.is_mlag_clone for mapping in context.mappings))

    def test_junos_l3_mlag_is_rejected_before_duplicate_addresses_are_generated(self):
        """ae0.0 的同一三层地址不能复制到两个独立物理口。"""
        topology, config_dir = conversion_fixture("junos_mlag")
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output"
            context = convert(topology, config_dir, output)
            self.assertTrue(context.has_errors)
            self.assertFalse((output / "configs").exists())
            self.assertFalse((output / "topology-adapted.xlsx").exists())
            self.assertIn("ae0", context.devices["JA"].document.render())
            self.assertIn("address 10.0.0.1/30;", context.devices["JA"].document.render())
            report = json.loads((output / "report.json").read_text(encoding="utf-8"))
            self.assertEqual(report["status"], "failed")
            self.assertTrue(any("ae0" in error and "设备 JA" in error for error in report["errors"]))
            self.assertFalse(any(mapping.is_mlag_clone for mapping in context.mappings))
            mappings = json.loads((output / "interface-mapping.json").read_text(encoding="utf-8"))
            removed_irb = next(item for item in mappings if item["source_interface"] == "irb.400")
            self.assertEqual(removed_irb["role"], "NON_NNI")
            self.assertEqual(removed_irb["action"], "remove")
            self.assertIsNone(removed_irb["target_interface"])

    def test_non_nni_duplicate_vlans_are_removed_without_reassignment(self):
        topology, config_dir = conversion_fixture("iosxr_duplicate_vlan")
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output"
            context = convert(topology, config_dir, output)
            self.assertFalse(context.has_errors)
            converted = (output / "configs" / "R1.cfg").read_text(encoding="utf-8")
            self.assertIn("interface GigabitEthernet0/0/0/0", converted)
            self.assertIn("ipv4 address 10.0.0.1 255.255.255.252", converted)
            self.assertNotIn("192.0.2.1", converted)
            self.assertNotIn("198.51.100.1", converted)
            self.assertNotIn("second-dot1q", converted)
            self.assertNotIn("interface GigabitEthernet0/0/0/5", converted)
            self.assertFalse(any(mapping.role == "UNI" for mapping in context.mappings))
            report = json.loads((output / "report.json").read_text(encoding="utf-8"))
            self.assertGreaterEqual(report["summary"]["group_conflict_count"], 1)
            self.assertEqual(report["group_conflicts"][0]["winner_source"], "explicit")

    def test_business_on_reserved_interfaces_is_removed_without_target(self):
        topology, source_config_dir = conversion_fixture("iosxr_bundle")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            config_dir = root / "configs"
            shutil.copytree(source_config_dir, config_dir)
            config_path = config_dir / "R1.cfg"
            source = config_path.read_text(encoding="utf-8")
            source = source.replace(
                "\nend\n",
                "\ninterface GigabitEthernet0/0/0/6\n"
                " ipv4 address 203.0.113.1 255.255.255.0\n!\n"
                "interface GigabitEthernet0/0/0/7\n"
                " ipv4 address 203.0.114.1 255.255.255.0\n!\nend\n",
            )
            config_path.write_text(source, encoding="utf-8")

            output = root / "output"
            context = convert(topology, config_dir, output)
            self.assertFalse(context.has_errors, context.errors)
            converted = (output / "configs" / "R1.cfg").read_text(encoding="utf-8")
            self.assertNotIn("interface GigabitEthernet0/0/0/6\n", converted)
            self.assertNotIn("interface GigabitEthernet0/0/0/7\n", converted)
            self.assertNotIn("203.0.113.1", converted)
            self.assertNotIn("203.0.114.1", converted)
            for number in (6, 7):
                self.assertTrue(any(
                    mapping.source_interface == f"GigabitEthernet0/0/0/{number}"
                    and mapping.role == "NON_NNI"
                    and mapping.action == "remove"
                    and mapping.target_interface is None
                    for mapping in context.devices["R1"].mappings
                ))
                self.assertFalse(any(
                    mapping.target_interface
                    and mapping.target_interface.startswith(f"GigabitEthernet0/0/0/{number}")
                    for mapping in context.devices["R1"].mappings
                ))
            report = json.loads((output / "report.json").read_text(encoding="utf-8"))
            adaptation = next(
                event for event in report["events"]
                if event["kind"] == "simulation-adaptation" and event["device"] == "R1"
            )
            self.assertNotIn("GigabitEthernet0/0/0/6", adaptation["target_interfaces"])
            self.assertNotIn("GigabitEthernet0/0/0/7", adaptation["target_interfaces"])

    def test_nni_capacity_excludes_three_reserved_interfaces(self):
        topology, config_dir = conversion_fixture("iosxr_mlag")
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            profile_path = root / "profiles.yaml"
            profile_path.write_text(
                "profiles:\n  cisco_iosxr:\n    interfaces:\n"
                "      - GigabitEthernet0/0/0/0\n"
                "      - GigabitEthernet0/0/0/5\n"
                "      - GigabitEthernet0/0/0/6\n"
                "      - GigabitEthernet0/0/0/7\n",
                encoding="utf-8",
            )
            context = convert(topology, config_dir, root / "output", profile_path)
            self.assertTrue(context.has_errors)
            self.assertTrue(any(
                "NNI 数量超过可用物理接口数 1" in error
                for error in context.errors
            ))

    def test_unresolved_applied_group_fails_conversion_before_mapping(self):
        topology, config_dir = conversion_fixture("junos_unresolved")
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output"
            context = convert(topology, config_dir, output)
            self.assertTrue(context.has_errors)
            self.assertFalse((output / "configs").exists())
            report = json.loads((output / "report.json").read_text(encoding="utf-8"))
            self.assertEqual(report["status"], "failed")
            self.assertTrue(any("group" in message for message in report["errors"]))

    def test_internal_invariant_failure_writes_report_then_reraises(self):
        """内部编程错误不被吞掉，但 CLI 仍能获得可诊断的失败报告。"""
        topology, config_dir = conversion_fixture("iosxr_bundle")

        class BrokenPipeline:
            def execute(self, context) -> None:
                raise InvariantViolation("测试内部不变量")

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output"
            with patch(
                "config_adaptor.adaptation.service.build_default_pipeline",
                return_value=BrokenPipeline(),
            ):
                with self.assertRaisesRegex(
                    InvariantViolation,
                    "测试内部不变量",
                ):
                    convert(topology, config_dir, output)

            self.assertFalse((output / "configs").exists())
            report = json.loads((output / "report.json").read_text(encoding="utf-8"))
            self.assertEqual(report["status"], "failed")
            self.assertTrue(
                any(
                    "InvariantViolation" in message
                    and "测试内部不变量" in message
                    for message in report["errors"]
                )
            )
            event = next(
                item for item in report["events"]
                if item["kind"] == "internal-error"
            )
            self.assertEqual(event["phase"], "核心转换流水线")
            self.assertEqual(event["exception_type"], "InvariantViolation")

    def test_render_invariant_failure_rewrites_report_as_failed(self):
        """输出阶段的结构错误不应留下成功报告或部分配置。"""
        topology, config_dir = conversion_fixture("junos_bundle")

        class CorruptingPipeline:
            def execute(self, context) -> None:
                context.devices["J1"].document.root.children = None

        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "output"
            with patch(
                "config_adaptor.adaptation.service.build_default_pipeline",
                return_value=CorruptingPipeline(),
            ):
                with self.assertRaisesRegex(
                    InvariantViolation,
                    "Junos 文档根节点",
                ):
                    convert(topology, config_dir, output)

            self.assertFalse((output / "configs").exists())
            report = json.loads((output / "report.json").read_text(encoding="utf-8"))
            self.assertEqual(report["status"], "failed")
            event = next(
                item for item in report["events"]
                if item["kind"] == "internal-error"
            )
            self.assertEqual(event["phase"], "输出结果")
            self.assertEqual(event["exception_type"], "InvariantViolation")


if __name__ == "__main__":
    unittest.main()
