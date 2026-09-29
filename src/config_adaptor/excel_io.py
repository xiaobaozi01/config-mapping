"""读取原始拓扑 Excel，并写出接口更新后的拓扑副本。"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .constants import (
    DEVICE_COLUMN_ALIASES,
    DEVICE_SHEET_ALIASES,
    LINK_COLUMN_ALIASES,
    LINK_SHEET_ALIASES,
)
from .models import Device, Link, TopologyWorkbook, Vendor


def _normalized(value: object) -> str:
    """统一表头和 Sheet 名的大小写、空白及下划线。"""
    return "".join(str(value or "").strip().lower().split()).replace("_", "")


def _find_sheet(sheet_names: Iterable[str], aliases: tuple[str, ...], kind: str) -> str:
    """按别名寻找工作表，并保留工作簿中的原始名称。"""
    lookup = {_normalized(name): name for name in sheet_names}
    for alias in aliases:
        if _normalized(alias) in lookup:
            return lookup[_normalized(alias)]
    raise ValueError(f"找不到{kind}工作表，允许名称: {', '.join(aliases)}")


def _header_map(sheet, aliases: dict[str, tuple[str, ...]], required: set[str]) -> dict[str, int]:
    """把逻辑字段名映射到 Excel 列号，同时检查必需列。"""
    raw = {
        _normalized(cell.value): cell.column
        for cell in sheet[1]
        if cell.value is not None and str(cell.value).strip()
    }
    result: dict[str, int] = {}
    for field, field_aliases in aliases.items():
        for alias in field_aliases:
            column = raw.get(_normalized(alias))
            if column:
                result[field] = column
                break
    missing = sorted(required - set(result))
    if missing:
        raise ValueError(f"工作表 {sheet.title} 缺少必需列: {', '.join(missing)}")
    return result


def load_topology(path: Path) -> TopologyWorkbook:
    """加载设备与链路，并转换成后续处理器使用的结构化模型。"""
    workbook = load_workbook(path)
    device_sheet_name = _find_sheet(workbook.sheetnames, DEVICE_SHEET_ALIASES, "设备列表")
    link_sheet_name = _find_sheet(workbook.sheetnames, LINK_SHEET_ALIASES, "链接表")
    device_sheet = workbook[device_sheet_name]
    link_sheet = workbook[link_sheet_name]
    device_headers = _header_map(device_sheet, DEVICE_COLUMN_ALIASES, {"name", "vendor"})
    link_headers = _header_map(
        link_sheet,
        LINK_COLUMN_ALIASES,
        {"a_device", "a_interface", "z_device", "z_interface"},
    )

    # 设备名是后续所有映射的主键，因此必须非空且全局唯一。
    devices: list[Device] = []
    seen: set[str] = set()
    for row in range(2, device_sheet.max_row + 1):
        name = str(device_sheet.cell(row, device_headers["name"]).value or "").strip()
        if not name:
            continue
        if name in seen:
            raise ValueError(f"设备列表存在重复设备名: {name}")
        seen.add(name)
        vendor_text = str(device_sheet.cell(row, device_headers["vendor"]).value or "").strip()
        config_file = None
        if "config_file" in device_headers:
            value = device_sheet.cell(row, device_headers["config_file"]).value
            config_file = str(value).strip() if value else None
        devices.append(Device(name=name, vendor=Vendor.parse(vendor_text), config_file=config_file, row=row))

    # 空行允许存在，但一旦某列有值就要求四个端点字段完整。
    links: list[Link] = []
    for row in range(2, link_sheet.max_row + 1):
        values = {
            field: str(link_sheet.cell(row, column).value or "").strip()
            for field, column in link_headers.items()
        }
        if not any(values.values()):
            continue
        if not all(values.get(field) for field in ("a_device", "a_interface", "z_device", "z_interface")):
            raise ValueError(f"链接表第 {row} 行端点信息不完整")
        links.append(Link(row=row, **values))

    return TopologyWorkbook(
        path=path,
        workbook=workbook,
        device_sheet=device_sheet_name,
        link_sheet=link_sheet_name,
        device_headers=device_headers,
        link_headers=link_headers,
        devices=devices,
        links=links,
    )


def write_adapted_topology(
    topology: TopologyWorkbook,
    output: Path,
    config_files: dict[str, str] | None = None,
) -> None:
    """复制原工作簿，改写有效链路并删除被扁平化的冗余行。"""

    # 重新加载源文件，确保内存中的解析对象和原始文件都不被修改。
    workbook = load_workbook(topology.path)
    sheet = workbook[topology.link_sheet]
    links_by_row = {link.row: link for link in topology.links}

    for row, link in links_by_row.items():
        if not link.active:
            continue
        sheet.cell(row, topology.link_headers["a_interface"], link.a_interface)
        sheet.cell(row, topology.link_headers["z_interface"], link.z_interface)

    # 必须从下往上删行，否则前面删除后会改变后续行号。
    for row in sorted((link.row for link in topology.links if not link.active), reverse=True):
        sheet.delete_rows(row, 1)

    if config_files and "config_file" in topology.device_headers:
        device_sheet = workbook[topology.device_sheet]
        name_column = topology.device_headers["name"]
        config_column = topology.device_headers["config_file"]
        for row in range(2, device_sheet.max_row + 1):
            name = str(device_sheet.cell(row, name_column).value or "").strip()
            if name in config_files:
                device_sheet.cell(row, config_column, config_files[name])

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
