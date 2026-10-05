"""读取原始拓扑 Excel，并写出接口更新后的拓扑副本。"""

from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from ..common.interface import interface_parent
from .models import ConversionContext, Device, InterfaceMapping, Link, TopologyWorkbook, Vendor


DEVICE_SHEET_ALIASES = ("设备列表", "devices", "device list", "device_list")
LINK_SHEET_ALIASES = ("链接表", "链路表", "links", "link list", "link_list")

DEVICE_COLUMN_ALIASES = {
    "name": ("设备名称", "设备名", "名称", "device", "device_name", "name", "hostname"),
    "vendor": ("厂商", "厂家", "vendor", "manufacturer"),
    "config_file": ("配置文件", "配置文件名", "config", "config_file", "configuration"),
}

LINK_COLUMN_ALIASES = {
    "a_device": ("a端设备", "源设备", "本端设备", "a_device", "device_a", "source_device"),
    "a_interface": ("a端接口", "源接口", "本端接口", "a_interface", "interface_a", "source_interface"),
    "z_device": ("z端设备", "对端设备", "目的设备", "z_device", "device_z", "target_device"),
    "z_interface": ("z端接口", "对端接口", "目的接口", "z_interface", "interface_z", "target_interface"),
}


def _normalized(value: object) -> str:
    """统一表头和 Sheet 名的大小写、空白及下划线。"""
    return "".join(str(value or "").strip().lower().split()).replace("_", "")


def _find_sheet(sheet_names: Iterable[str], aliases: tuple[str, ...], kind: str) -> str:
    """按别名寻找工作表，并保留工作簿中的原始名称。"""
    lookup = {_normalized(name): name for name in sheet_names}
    for alias in aliases:
        if _normalized(alias) in lookup:
            return lookup[_normalized(alias)]
    raise ValueError(f"找不到{kind}工作表，允许名称: {', '.join(aliases)}")


def _header_map(sheet, aliases: dict[str, tuple[str, ...]], required: set[str]) -> dict[str, int]:
    """把逻辑字段名映射到 Excel 列号，同时检查必需列。"""
    raw = {
        _normalized(cell.value): cell.column
        for cell in sheet[1]
        if cell.value is not None and str(cell.value).strip()
    }
    result: dict[str, int] = {}
    for field, field_aliases in aliases.items():
        for alias in field_aliases:
            column = raw.get(_normalized(alias))
            if column:
                result[field] = column
                break
    missing = sorted(required - set(result))
    if missing:
        raise ValueError(f"工作表 {sheet.title} 缺少必需列: {', '.join(missing)}")
    return result


def load_topology(path: Path) -> TopologyWorkbook:
    """加载设备与链路，并转换成后续处理器使用的结构化模型。"""
    workbook = load_workbook(path)
    device_sheet_name = _find_sheet(workbook.sheetnames, DEVICE_SHEET_ALIASES, "设备列表")
    link_sheet_name = _find_sheet(workbook.sheetnames, LINK_SHEET_ALIASES, "链接表")
    device_sheet = workbook[device_sheet_name]
    link_sheet = workbook[link_sheet_name]
    device_headers = _header_map(device_sheet, DEVICE_COLUMN_ALIASES, {"name", "vendor"})
    link_headers = _header_map(
        link_sheet,
        LINK_COLUMN_ALIASES,
        {"a_device", "a_interface", "z_device", "z_interface"},
    )

    # 设备名是后续所有映射的主键，因此必须非空且全局唯一。
    devices: list[Device] = []
    seen: set[str] = set()
    for row in range(2, device_sheet.max_row + 1):
        name = str(device_sheet.cell(row, device_headers["name"]).value or "").strip()
        if not name:
            continue
        if name in seen:
            raise ValueError(f"设备列表存在重复设备名: {name}")
        seen.add(name)
        vendor_text = str(device_sheet.cell(row, device_headers["vendor"]).value or "").strip()
        config_file = None
        if "config_file" in device_headers:
            value = device_sheet.cell(row, device_headers["config_file"]).value
            config_file = str(value).strip() if value else None
        devices.append(Device(name=name, vendor=Vendor.parse(vendor_text), config_file=config_file, row=row))

    # 空行允许存在，但一旦某列有值就要求四个端点字段完整。
    links: list[Link] = []
    for row in range(2, link_sheet.max_row + 1):
        values = {
            field: str(link_sheet.cell(row, column).value or "").strip()
            for field, column in link_headers.items()
        }
        if not any(values.values()):
            continue
        if not all(values.get(field) for field in ("a_device", "a_interface", "z_device", "z_interface")):
            raise ValueError(f"链接表第 {row} 行端点信息不完整")
        links.append(Link(row=row, **values))

    return TopologyWorkbook(
        path=path,
        workbook=workbook,
        device_sheet=device_sheet_name,
        link_sheet=link_sheet_name,
        device_headers=device_headers,
        link_headers=link_headers,
        devices=devices,
        links=links,
    )


def write_adapted_topology(
    topology: TopologyWorkbook,
    output: Path,
    config_files: dict[str, str] | None = None,
) -> None:
    """复制原工作簿，改写有效链路并删除被扁平化的冗余行。"""

    # 重新加载源文件，确保内存中的解析对象和原始文件都不被修改。
    workbook = load_workbook(topology.path)
    sheet = workbook[topology.link_sheet]
    links_by_row = {link.row: link for link in topology.links}

    for row, link in links_by_row.items():
        if not link.active:
            continue
        sheet.cell(row, topology.link_headers["a_interface"], link.a_interface)
        sheet.cell(row, topology.link_headers["z_interface"], link.z_interface)

    # 必须从下往上删行，否则前面删除后会改变后续行号。
    for row in sorted((link.row for link in topology.links if not link.active), reverse=True):
        sheet.delete_rows(row, 1)

    if config_files and "config_file" in topology.device_headers:
        device_sheet = workbook[topology.device_sheet]
        name_column = topology.device_headers["name"]
        config_column = topology.device_headers["config_file"]
        for row in range(2, device_sheet.max_row + 1):
            name = str(device_sheet.cell(row, name_column).value or "").strip()
            if name in config_files:
                device_sheet.cell(row, config_column, config_files[name])

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
class TopologyPreflightHandler:
    """在配置改写前划定本次转换能够安全处理的拓扑链路范围。

    该处理器检查每条链路的两端设备是否存在于拓扑设备表，以及是否已加载为
    当前版本支持的设备；缺失设备的链路记为错误，不支持厂商的链路则记为警告
    并跳过。预检必须位于流水线最前面，否则无效链路可能消耗 NNI 目标端口、
    触发不存在设备的访问，甚至让后续 UNI 阶段错误迁移原本属于该链路的接口。
    """

    def process(self, context: ConversionContext) -> None:
        """校验全部链路端点，停用不可处理的链路并记录诊断及汇总事件。

        端点设备未出现在设备表时说明输入拓扑不完整，因此写入错误，让流水线在
        本阶段后停止；设备存在但厂商暂不支持时只写警告和 ``skip-link`` 事件，
        允许其余链路继续转换。两类链路都会设置 ``active=False``，因为下游阶段
        只应消费已通过预检的链路，同时会为可识别端点补记跳过映射，保留其 NNI
        身份。最后记录活动/跳过数量，供报告解释哪些拓扑数据实际参与了转换。
        """
        topology_devices = {device.name for device in context.topology.devices}
        for link in context.topology.links:
            missing = [name for name, _ in link.endpoints() if name not in topology_devices]
            if missing:
                link.active = False
                link.skip_reason = f"端点设备未出现在设备列表: {', '.join(missing)}"
                context.errors.append(f"链接表第 {link.row} 行：{link.skip_reason}")
                self._record_skipped_endpoints(context, link, link.skip_reason)
                continue
            if all(device_name in context.devices for device_name, _ in link.endpoints()):
                continue
            link.active = False
            link.skip_reason = "端点包含不在首版范围内的设备（可能为华为或未知厂商）"
            self._record_skipped_endpoints(context, link, link.skip_reason)
            context.warnings.append(f"链接表第 {link.row} 行已跳过：{link.skip_reason}")
            context.add_event(
                "skip-link",
                f"跳过链接表第 {link.row} 行",
                row=link.row,
                reason=link.skip_reason,
            )

        context.add_event(
            "topology-preflight",
            "拓扑链路预检完成",
            active_links=sum(1 for link in context.topology.links if link.active),
            skipped_links=sum(1 for link in context.topology.links if not link.active),
        )

    @staticmethod
    def _record_skipped_endpoints(
        context: ConversionContext,
        link: Link,
        reason: str | None,
    ) -> None:
        """为跳过链路中仍可识别的端点登记 NNI 保留映射。

        方法同时记录规范化物理父口及其可能所属的 Bundle/ae 逻辑口，但不生成
        目标接口。这样 UNI 候选筛选仍会排除这些源口，避免因为链路被停用就把
        原 NNI 业务误判成 UNI 并迁移；不存在或不支持的设备端点则安全忽略。
        """
        for device_name, raw_interface in link.endpoints():
            device = context.devices.get(device_name)
            if not device:
                continue
            source = interface_parent(device.document.resolve_interface(raw_interface))
            logical = device.document.bundle_members().get(source)
            for reserved in dict.fromkeys((source, logical)):
                if not reserved:
                    continue
                device.mappings.append(
                    InterfaceMapping(
                        device=device_name,
                        source_interface=reserved,
                        role="NNI",
                        action="skip",
                        target_interface=None,
                        link_rows=[link.row],
                        reason=reason,
                    )
                )
